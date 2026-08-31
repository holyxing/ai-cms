#!/usr/bin/env python3
"""生成 AI-CMS《产品使用手册》与《HY 标签使用手册》Word 文档。"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs"
IMG = OUT / "images"
BLUE = RGBColor(0x25, 0x63, 0xEB)
SLATE = RGBColor(0x0F, 0x17, 0x2A)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_run_font(run, name="微软雅黑", size=11, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def setup_doc(title: str) -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run(f"AI-CMS  ·  {title}")
    set_run_font(r, size=9, color=MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run("AI-CMS  ·  使用手册  ·  请勿将 API Key 写入模板或文档截图")
    set_run_font(r, size=9, color=MUTED)
    return doc


def add_cover(doc: Document, title: str, subtitle: str):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("AI-CMS")
    set_run_font(r, size=16, bold=True, color=BLUE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    set_run_font(r, size=28, bold=True, color=SLATE)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    set_run_font(r, size=12, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("版本 0.1 · 2026-08-24")
    set_run_font(r, size=11, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("自托管多站点内容管理与静态发布平台")
    set_run_font(r, size=11, color=MUTED)
    doc.add_page_break()


def h(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, size={1: 18, 2: 14, 3: 12}.get(level, 11), bold=True, color=BLUE if level <= 2 else SLATE)


def para(doc, text, size=11, bold=False, color=SLATE):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    return p


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    r = p.add_run(text)
    set_run_font(r, size=11)
    p.paragraph_format.left_indent = Cm(0.75 * (level + 1))
    return p


def numbered(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.clear()
    r = p.add_run(text)
    set_run_font(r, size=11)
    return p


def note(doc, text, kind="注意"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(f"【{kind}】")
    set_run_font(r, size=11, bold=True, color=BLUE)
    r2 = p.add_run(text)
    set_run_font(r2, size=11, color=SLATE)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    for i, hd in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        r = p.add_run(hd)
        set_run_font(r, size=10, bold=True, color=WHITE)
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "2563EB")
        tcPr.append(shd)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            r = p.add_run(str(val))
            set_run_font(r, size=10)
    doc.add_paragraph()


def code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text.strip("\n"))
    set_run_font(r, name="Consolas", size=9, color=SLATE)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F1F5F9")
    pPr.append(shd)


def add_image(doc, name, caption=None, width_cm=15.5):
    path = IMG / name
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    if caption:
        c = doc.add_paragraph()
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.add_run(caption)
        set_run_font(r, size=9, color=MUTED)


def add_steps(doc, items: list[str]):
    for i, s in enumerate(items, 1):
        numbered(doc, s)


# ---------- Markdown → docx ----------
_IMG_MD = re.compile(r"^!\[([^\]]*)\]\((?:images/)?([^)]+)\)$")


def md_to_docx(doc: Document, md_path: Path, skip_first_h1: bool = False):
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    i = 0
    # skip yaml
    if lines and lines[0].strip() == "---":
        i = 1
        while i < len(lines) and lines[i].strip() != "---":
            i += 1
        i += 1
    in_code = False
    code_buf: list[str] = []
    table_buf: list[str] = []
    skipped_h1 = False

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        rows = []
        for tl in table_buf:
            if re.match(r"^\|?\s*-+", tl):
                continue
            cells = [c.strip() for c in tl.strip().strip("|").split("|")]
            cells = [re.sub(r"\*\*(.*?)\*\*", r"\1", c) for c in cells]
            rows.append(cells)
        table_buf = []
        if not rows:
            return
        headers, data = rows[0], rows[1:]
        table(doc, headers, data)

    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if in_code:
                code_block(doc, "\n".join(code_buf))
                code_buf = []
                in_code = False
            else:
                flush_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue
        if line.strip().startswith("|"):
            table_buf.append(line)
            i += 1
            continue
        flush_table()
        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        img_m = _IMG_MD.match(line.strip())
        if img_m:
            add_image(doc, img_m.group(2).strip(), img_m.group(1).strip() or None)
            i += 1
            continue
        if line.startswith("# "):
            if skip_first_h1 and not skipped_h1:
                skipped_h1 = True
            else:
                h(doc, line[2:].strip(), 1)
        elif re.match(r"^\d+\.\s+", line.strip()):
            numbered(doc, re.sub(r"\*\*(.*?)\*\*", r"\1", re.sub(r"^\d+\.\s+", "", line.strip())))
        elif line.startswith("## "):
            h(doc, line[3:].strip(), 2)
        elif line.startswith("### "):
            h(doc, line[4:].strip(), 3)
        elif line.startswith("> "):
            note(doc, re.sub(r"\*\*(.*?)\*\*", r"\1", line[2:].strip()), "说明")
        elif line.startswith("- "):
            bullet(doc, re.sub(r"\*\*(.*?)\*\*", r"\1", line[2:].strip()))
        else:
            para(doc, re.sub(r"\*\*(.*?)\*\*", r"\1", line.strip()))
        i += 1
    flush_table()


def build_tag_handbook() -> Path:
    doc = setup_doc("HY 标签使用手册")
    add_cover(
        doc,
        "HY 标签使用手册",
        "模板占位符逐条说明 · 给会一点 HTML 的编辑与运营",
    )
    h(doc, "阅读前请先知道", 1)
    para(doc, "HY_ 标签不是普通 HTML，而是 AI-CMS 在「发布生成静态页」时替换成真实内容的占位符。你在模板里写标签，访客看到的是替换后的网站。")
    para(doc, "本手册依据当前渲染器白名单与《HY_标签逐条使用手册》整理。每个标签都有「模板示例」和「渲染后 HTML」对照。模板编辑器中输入 HY_ 可自动补全。")
    add_image(doc, "layouts.png", "图：模板编辑器（HTML 源码 + HY_ 标签）")
    md_to_docx(doc, OUT / "HY_标签逐条使用手册.md")
    h(doc, "渲染器补充标签（手册原文可能未单独成节）", 1)
    para(doc, "以下标签已在系统中实现，写模板时可以使用。")
    table(doc, ["标签", "说明"], [
        ["HY_SITE_LINK", "站点内页链接，按页面深度自动加 ../"],
        ["HY_ITEM_BANNER", "列表项 Banner 大图 URL（需在 HY_CONTENTS 内）"],
        ["HY_ITEM_SUBTITLE", "列表项副标题"],
        ["HY_ITEM_CAT_SLUG", "列表项所属栏目 slug"],
        ["HY_CAT_PARENT_URL", "父栏目 URL"],
        ["HY_CAT_ITEM_CURRENT_CLASS", "栏目循环中当前栏目的 CSS class"],
        ["HY_CONTENT_SUMMARY / HY_CONTENT_EXCERPT", "详情摘要"],
        ["HY_CONTENT_PUBLISH_DATE", "发布日期"],
        ["HY_CONTENT_READ_TIME", "预估阅读时长"],
        ["HY_TEMPLATE", "推荐写法：<HY_TEMPLATE _code=\"页头code\" /> 或 code=\"...\"（以实现为准）"],
    ])
    h(doc, "给新手的最小模板", 1)
    para(doc, "栏目列表页只要会这几行就能出列表：")
    code_block(doc, """<h1><HY_CAT_NAME /></h1>
<HY_CONTENTS _limit="12" _order="newest">
  <a href="<HY_ITEM_URL />"><HY_ITEM_TITLE /></a>
</HY_CONTENTS>
<HY_CONTENTS_EMPTY><p>本栏目暂无文章</p></HY_CONTENTS_EMPTY>
<HY_CONTENTS_PAGINATION />""")
    para(doc, "文章详情页最小结构：")
    code_block(doc, """<h1><HY_CONTENT_TITLE /></h1>
<p><HY_CONTENT_DATE /> · <HY_CONTENT_AUTHOR /></p>
<div><HY_CONTENT_BODY /></div>""")
    note(doc, "列表循环里用 HY_ITEM_*，详情页用 HY_CONTENT_*。两者不能混用。", "必记")
    out = OUT / "AI-CMS-HY标签使用手册.docx"
    doc.save(out)
    return out


def build_product_handbook() -> Path:
    doc = setup_doc("产品使用手册")
    add_cover(
        doc,
        "产品使用手册",
        "按后台每一个页面与按钮对照 · 逐步操作",
    )
    note(
        doc,
        "上一版手册偏「概念导读」，漏写了栏目 Excel 导入、文章版本恢复、栏目批量操作、站点资源、发布回滚与文章回退的区别等。本版按 frontend 路由与页面控件重写，功能以界面为准。",
        "修订说明",
    )
    md_to_docx(doc, OUT / "AI-CMS-产品使用手册.md", skip_first_h1=True)
    out = OUT / "AI-CMS-产品使用手册.docx"
    doc.save(out)
    return out


def main():
    p1 = build_product_handbook()
    p2 = build_tag_handbook()
    print(p1)
    print(p2)


if __name__ == "__main__":
    main()
